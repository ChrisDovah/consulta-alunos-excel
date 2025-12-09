import pandas as pd
from datetime import datetime
import os
import unicodedata

def remover_acentos(texto):
    """Remove acentos e caracteres especiais de um texto"""
    if pd.isna(texto):
        return ""
    texto = str(texto)
    # Normaliza o texto (NFD = Canonical Decomposition)
    # Depois remove os caracteres de combinação (acentos)
    nfd = unicodedata.normalize('NFD', texto)
    texto_sem_acento = ''.join(char for char in nfd if unicodedata.category(char) != 'Mn')
    return texto_sem_acento

def carregar_planilha(caminho_arquivo):
    """Carrega a planilha Excel e retorna o DataFrame"""
    try:
        # Tentar carregar sem especificar o nome da aba (pega a primeira)
        df = pd.read_excel(caminho_arquivo, sheet_name=0)
        
        # Verificar se tem pelo menos 3 colunas
        if len(df.columns) < 3:
            print(f"❌ Erro: A planilha deve ter pelo menos 3 colunas.")
            return None
        
        # Renomear as 3 primeiras colunas para facilitar o trabalho
        colunas_antigas = df.columns.tolist()
        df = df.rename(columns={
            colunas_antigas[0]: 'Data_Hora',
            colunas_antigas[1]: 'Pontuacao',
            colunas_antigas[2]: 'Nome_Completo'
        })
        
        print(f"✅ Colunas identificadas:")
        print(f"   Coluna A (Data/Hora): '{colunas_antigas[0]}'")
        print(f"   Coluna B (Pontuação): '{colunas_antigas[1]}'")
        print(f"   Coluna C (Nome): '{colunas_antigas[2]}'")
        
        return df
    except FileNotFoundError:
        print(f"❌ Erro: Arquivo '{caminho_arquivo}' não encontrado.")
        return None
    except Exception as e:
        print(f"❌ Erro ao carregar planilha: {e}")
        return None

def extrair_nota(pontuacao_str):
    """Extrai a nota do formato '8 / 10' e retorna apenas o valor numérico"""
    try:
        if pd.isna(pontuacao_str):
            return None
        nota_str = str(pontuacao_str).split('/')[0].strip()
        return float(nota_str)
    except:
        return None

def verificar_situacao(nota):
    """Verifica se o aluno foi aprovado ou reprovado"""
    if nota is None:
        return None
    try:
        if float(nota) >= 5:
            return "APROVADO"
        else:
            return "REPROVADO"
    except:
        return None

def formatar_data(data):
    """Formata a data para um padrão legível"""
    try:
        if pd.isna(data):
            return "Sem data"
        if isinstance(data, str):
            return data
        return data.strftime('%d/%m/%Y %H:%M:%S')
    except:
        return "Sem data"

def consultar_alunos(df, lista_nomes, debug=False):
    """Consulta cada aluno na planilha e retorna os resultados"""
    resultados = []
    
    # Criar coluna normalizada uma única vez (mais eficiente)
    df['Nome_Normalizado'] = df['Nome_Completo'].apply(
        lambda x: remover_acentos(str(x)).upper().strip() if pd.notna(x) else ""
    )
    
    for nome in lista_nomes:
        nome_original = nome.strip()
        # Normalizar: remover acentos, converter para maiúsculas e remover espaços extras
        nome_normalizado = remover_acentos(nome_original).upper().strip()
        
        if debug:
            print(f"\n🔍 Buscando: '{nome_original}'")
            print(f"   Normalizado: '{nome_normalizado}'")
        
        # Buscar aluno usando os nomes normalizados
        aluno_encontrado = df[df['Nome_Normalizado'] == nome_normalizado]
        
        if debug and aluno_encontrado.empty:
            # Mostrar nomes similares para debug
            similares = df[df['Nome_Normalizado'].str.contains(nome_normalizado.split()[0], na=False)]
            if not similares.empty:
                print(f"   ⚠️ Nomes similares encontrados:")
                for idx, row in similares.head(3).iterrows():
                    print(f"      - '{row['Nome_Completo']}' → '{row['Nome_Normalizado']}'")
        
        if not aluno_encontrado.empty:
            # Se houver duplicatas, pegar a última ocorrência
            ultima_ocorrencia = aluno_encontrado.iloc[-1]
            nota = extrair_nota(ultima_ocorrencia['Pontuacao'])
            
            resultados.append({
                'nome': nome_original,
                'status': 'CONSTA',
                'nota': nota,
                'situacao': verificar_situacao(nota),
                'pontuacao_completa': ultima_ocorrencia['Pontuacao'],
                'data_hora': formatar_data(ultima_ocorrencia['Data_Hora'])
            })
        else:
            resultados.append({
                'nome': nome_original,
                'status': 'NÃO CONSTA',
                'nota': None,
                'situacao': None,
                'pontuacao_completa': '-',
                'data_hora': '-'
            })
    
    return resultados

def gerar_relatorio_excel(resultados, arquivo_saida='relatorio_alunos.xlsx'):
    """Gera um relatório em formato Excel"""
    # Preparar dados para DataFrame
    dados = []
    for resultado in resultados:
        dados.append({
            'Nome do Aluno': resultado['nome'],
            'Status': resultado['status'],
            'Pontuação': resultado['pontuacao_completa'] if resultado['status'] == 'CONSTA' else '-',
            'Situação': resultado['situacao'] if resultado['situacao'] else '-',
            'Data/Hora': resultado['data_hora']
        })
    
    # Criar DataFrame
    df = pd.DataFrame(dados)
    
    # Criar arquivo Excel com formatação
    with pd.ExcelWriter(arquivo_saida, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Relatório de Alunos', index=False)
        
        # Obter a planilha para formatação
        workbook = writer.book
        worksheet = writer.sheets['Relatório de Alunos']
        
        # Ajustar largura das colunas
        worksheet.column_dimensions['A'].width = 40  # Nome
        worksheet.column_dimensions['B'].width = 15  # Status
        worksheet.column_dimensions['C'].width = 15  # Pontuação
        worksheet.column_dimensions['D'].width = 15  # Situação
        worksheet.column_dimensions['E'].width = 20  # Data/Hora
        
        # Formatar cabeçalho (negrito e cor de fundo)
        from openpyxl.styles import Font, PatternFill, Alignment
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Centralizar colunas B, C, D, E
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for idx, cell in enumerate(row):
                if idx > 0:  # Pula a primeira coluna (nome)
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    
                # Colorir linha de acordo com situação
                if idx == 3:  # Coluna Situação
                    if cell.value == 'APROVADO':
                        for c in row:
                            c.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                    elif cell.value == 'REPROVADO':
                        for c in row:
                            c.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    
    print(f"\n✅ Relatório Excel gerado com sucesso: {arquivo_saida}")

def gerar_relatorio_txt(resultados, arquivo_saida='relatorio_alunos.txt'):
    """Gera um relatório em formato texto"""
    with open(arquivo_saida, 'w', encoding='utf-8') as f:
        f.write("=" * 100 + "\n")
        f.write("RELATÓRIO DE CONSULTA DE ALUNOS\n")
        f.write("=" * 100 + "\n")
        f.write(f"Data de geração: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        f.write(f"Total de alunos consultados: {len(resultados)}\n")
        f.write("=" * 100 + "\n\n")
        
        # Contar estatísticas
        consta = sum(1 for r in resultados if r['status'] == 'CONSTA')
        nao_consta = len(resultados) - consta
        
        f.write(f"📊 RESUMO:\n")
        f.write(f"   ✓ Alunos encontrados: {consta}\n")
        f.write(f"   ✗ Alunos não encontrados: {nao_consta}\n")
        f.write("\n" + "=" * 100 + "\n\n")
        
        # Detalhamento de cada aluno
        f.write("DETALHAMENTO POR ALUNO:\n")
        f.write("-" * 100 + "\n\n")
        
        for i, resultado in enumerate(resultados, 1):
            f.write(f"{i}. {resultado['nome']}\n")
            f.write(f"   Status: {resultado['status']}\n")
            
            if resultado['status'] == 'CONSTA':
                # Mostrar apenas a pontuação completa (ex: 8 / 10)
                f.write(f"   Pontuação: {resultado['pontuacao_completa']}\n")
                # Mostrar situação (APROVADO/REPROVADO)
                if resultado['situacao']:
                    f.write(f"   Situação: {resultado['situacao']}\n")
                # Mostrar data/hora
                f.write(f"   Data/Hora: {resultado['data_hora']}\n")
            
            f.write("\n")
        
        f.write("=" * 100 + "\n")
        f.write("FIM DO RELATÓRIO\n")
        f.write("=" * 100 + "\n")
    
    print(f"\n✅ Relatório gerado com sucesso: {arquivo_saida}")

def carregar_lista_nomes(caminho_arquivo=None):
    """Carrega lista de nomes de um arquivo txt ou permite entrada manual"""
    if caminho_arquivo and os.path.exists(caminho_arquivo):
        with open(caminho_arquivo, 'r', encoding='utf-8') as f:
            nomes = [linha.strip() for linha in f.readlines() if linha.strip()]
        return nomes
    else:
        print("\n📝 Digite os nomes dos alunos (um por linha).")
        print("Quando terminar, deixe uma linha em branco e pressione Enter:\n")
        nomes = []
        while True:
            nome = input().strip()
            if not nome:
                break
            nomes.append(nome)
        return nomes

def main():
    print("=" * 100)
    print("SISTEMA DE CONSULTA DE ALUNOS EM PLANILHA EXCEL")
    print("=" * 100)
    
    # 1. Solicitar caminho da planilha Excel
    print("\n📂 Digite o caminho completo da planilha Excel:")
    print("   Exemplo: C:\\Users\\SeuNome\\Documents\\planilha.xlsx")
    caminho_excel = input("   Caminho: ").strip().strip('"')
    
    # Carregar planilha
    df = carregar_planilha(caminho_excel)
    if df is None:
        return
    
    print(f"\n✅ Planilha carregada com sucesso! Total de registros: {len(df)}")
    
    # 2. Solicitar lista de nomes
    print("\n" + "=" * 100)
    print("OPÇÕES DE ENTRADA DOS NOMES:")
    print("1 - Digitar os nomes manualmente")
    print("2 - Carregar de um arquivo .txt")
    opcao = input("Escolha uma opção (1 ou 2): ").strip()
    
    if opcao == '2':
        print("\n📂 Digite o caminho do arquivo .txt com os nomes:")
        caminho_txt = input("   Caminho: ").strip().strip('"')
        lista_nomes = carregar_lista_nomes(caminho_txt)
    else:
        lista_nomes = carregar_lista_nomes()
    
    if not lista_nomes:
        print("\n❌ Nenhum nome foi fornecido.")
        return
    
    print(f"\n✅ {len(lista_nomes)} nome(s) carregado(s).")
    
    # 3. Consultar alunos
    print("\n🔍 Processando consultas...")
    print("\n⚙️ Deseja ativar modo DEBUG para ver detalhes da busca? (s/n): ", end="")
    debug_mode = input().strip().lower() == 's'
    
    resultados = consultar_alunos(df, lista_nomes, debug=debug_mode)
    
    # 4. Escolher formato de saída
    print("\n" + "=" * 100)
    print("FORMATO DO RELATÓRIO:")
    print("1 - Arquivo TXT (texto)")
    print("2 - Arquivo Excel (.xlsx)")
    print("3 - Ambos (TXT e Excel)")
    opcao_formato = input("Escolha o formato (1, 2 ou 3): ").strip()
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    
    # 5. Gerar relatório conforme escolha
    print("\n📄 Gerando relatório(s)...")
    
    if opcao_formato in ['1', '3']:
        nome_relatorio_txt = f"relatorio_alunos_{timestamp}.txt"
        gerar_relatorio_txt(resultados, nome_relatorio_txt)
    
    if opcao_formato in ['2', '3']:
        nome_relatorio_excel = f"relatorio_alunos_{timestamp}.xlsx"
        gerar_relatorio_excel(resultados, nome_relatorio_excel)
    
    # 6. Exibir resumo no console
    print("\n" + "=" * 100)
    print("RESUMO DA CONSULTA:")
    print("=" * 100)
    consta = sum(1 for r in resultados if r['status'] == 'CONSTA')
    print(f"✓ Alunos encontrados: {consta}/{len(resultados)}")
    print(f"✗ Alunos não encontrados: {len(resultados) - consta}/{len(resultados)}")
    print("=" * 100)

if __name__ == "__main__":
    main()